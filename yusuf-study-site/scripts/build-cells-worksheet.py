#!/usr/bin/env python3
"""Build the optional original Cell Theory worksheet and separate key.

Requires python-docx. Run with the bundled document runtime (or an environment
where python-docx is installed): python scripts/build-cells-worksheet.py.
Outputs contain native editable Word text controls and no learner responses.
After changes, render both documents and inspect every page before publishing.
"""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public/quizzes/cells/concept-1'
NAMES = ['Robert Hooke', 'Anton van Leeuwenhoek', 'Matthias Schleiden', 'Theodor Schwann', 'Rudolf Virchow']
YEARS = [1665, 1673, 1838, 1839, 1858]
DISCOVERIES = [
    'Observed cork compartments and named them cells. The compartments were dead plant-cell walls.',
    'Began reporting microscope observations. Later observations revealed microscopic living organisms.',
    'Concluded that all plants are made of cells.',
    'Concluded that animals are made of cells. His work and Schleiden’s supported the first two principles.',
    'Helped establish that new cells come from existing cells. His book helped spread the idea.',
]
MC = [
    ('Which scientist observed cork and named the structures cells?', ['Robert Hooke', 'Theodor Schwann', 'Rudolf Virchow'], 'A'),
    ('Which scientist concluded that plants are made of cells?', ['Anton van Leeuwenhoek', 'Matthias Schleiden', 'Robert Hooke'], 'B'),
    ('Which two scientists helped establish the first two parts of Cell Theory?', ['Hooke and Virchow', 'Leeuwenhoek and Hooke', 'Schleiden and Schwann'], 'C'),
    ('Which scientist helped establish the third part of Cell Theory?', ['Rudolf Virchow', 'Robert Hooke', 'Matthias Schleiden'], 'A'),
    ('Which discovery came first in this story?', ['Animals are made of cells', 'Cork contains tiny compartments', 'New cells come from existing cells'], 'B'),
]
TYPED = [
    ('What are the three parts of Cell Theory?', 'All living organisms are made of one or more cells. The cell is the basic unit of structure and function. All cells come from pre-existing cells.'),
    ('How did Hooke’s discovery lead to further investigation?', 'Hooke named tiny compartments in cork. This raised questions about whether microscopic structures and living things existed in other materials. Later observations expanded the evidence.'),
    ('Why were Schleiden and Schwann important together?', 'Schleiden concluded that plants are made of cells, and Schwann concluded that animals are too. Their work helped establish cells as the basic unit of living things.'),
    ('How did Virchow contribute to Cell Theory?', 'Virchow helped establish and spread the idea that new cells come from cells that already exist. This is the third principle.'),
    ('Why was the microscope important?', 'Microscopes made tiny structures and organisms visible. Scientists could gather evidence that they could not see with the naked eye.'),
]
FIELD_ID = 8000


def set_style(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = bold
    style.paragraph_format.space_after = Pt(7)
    # Some bundled base templates carry a blue rule in the Title style.
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)


def document(title, intro):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)
    set_style(doc.styles['Normal'], 11.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    set_style(doc.styles['Title'], 25)
    set_style(doc.styles['Heading 1'], 17, True)
    set_style(doc.styles['Heading 2'], 13, True)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.core_properties.author = 'Study Club'
    doc.core_properties.last_modified_by = 'Study Club'
    doc.core_properties.title = title
    doc.core_properties.subject = 'Science Cells Concept 1'
    doc.core_properties.comments = ''
    doc.core_properties.created = datetime(2026, 10, 19, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 10, 19, tzinfo=timezone.utc)
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(intro)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run('Science  |  Cells  |  Concept 1')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string('555555')
    return doc


def field(paragraph, tag, placeholder='Type here', lines=1):
    """Append an unlocked plain-text content control, supporting multiline input."""
    global FIELD_ID
    FIELD_ID += 1
    sdt = OxmlElement('w:sdt')
    properties = OxmlElement('w:sdtPr')
    for name, value in [('alias', tag.replace('_', ' ')), ('tag', tag), ('id', str(FIELD_ID))]:
        element = OxmlElement('w:' + name)
        element.set(qn('w:val'), value)
        properties.append(element)
    text = OxmlElement('w:text')
    text.set(qn('w:multiLine'), '1')
    properties.append(text)
    sdt.append(properties)
    content = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r')
    run_properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '4E625D')
    run_properties.append(color)
    run.append(run_properties)
    value = OxmlElement('w:t')
    value.text = '[' + placeholder + ']'
    run.append(value)
    for _ in range(lines - 1):
        run.append(OxmlElement('w:br'))
        blank = OxmlElement('w:t')
        blank.set(qn('xml:space'), 'preserve')
        blank.text = ' '
        run.append(blank)
    content.append(run)
    sdt.append(content)
    paragraph._p.append(sdt)


def response(doc, tag, prefix='', lines=1):
    p = doc.add_paragraph(prefix)
    p.paragraph_format.keep_together = True
    field(p, tag, 'Click or tap here to type' if lines > 1 else 'Type here', lines)
    return p


def page(doc, heading):
    doc.add_page_break()
    doc.add_heading(heading, 1)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    properties = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), '4')
        edge.set(qn('w:color'), 'D9D9D9')
        borders.append(edge)
    properties.append(borders)
    for index, row_values in enumerate([headers] + rows):
        row = t.rows[0] if index == 0 else t.add_row()
        row_prop = row._tr.get_or_add_trPr()
        if index == 0:
            row_prop.append(OxmlElement('w:tblHeader'))
        row_prop.append(OxmlElement('w:cantSplit'))
        for cell, value, width in zip(row.cells, row_values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            tc = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                margin = OxmlElement('w:' + side)
                margin.set(qn('w:w'), '100')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            tc.append(margins)
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'DDECE7')
                tc.append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.bold = index == 0
    return t


def build_worksheet():
    doc = document('Cell Theory and Scientists Practice Worksheet', 'Follow the discoveries that helped scientists explain cells. Work through one activity at a time. Click a bracketed field to type, or print and write. Save a copy to keep your answers.')
    doc.add_heading('Timeline reference', 1)
    table(doc, ['Year', 'Scientist', 'How the story developed'], list(zip(YEARS, NAMES, DISCOVERIES)), [.62, 1.5, 4.88])
    doc.add_heading('A Complete the three parts', 1)
    doc.add_paragraph('Word bank: cells  |  basic unit  |  pre-existing cells')
    for i, (before, after) in enumerate([
        ('1. All living organisms are made of one or more ', '.'),
        ('2. The cell is the ', ' of structure and function in living organisms.'),
        ('3. All cells come from ', '.'),
    ]):
        p = doc.add_paragraph(before)
        field(p, f'principle_{i+1}')
        p.add_run(after)
    doc.add_paragraph('One organism can have one cell or many cells. Pre-existing means already there.')

    page(doc, 'B Match each scientist to a discovery')
    doc.add_paragraph('Write the correct letter next to each scientist. Use each letter once.')
    for i, name in enumerate(NAMES):
        response(doc, f'match_{i+1}', f'{i+1}. {name}: ')
    for text in ['A  Concluded that animals are made of cells.', 'B  Helped establish that new cells come from existing cells.', 'C  Observed cork compartments and named them cells.', 'D  Concluded that all plants are made of cells.', 'E  Observed tiny living organisms called animalcules.']:
        doc.add_paragraph(text)
    doc.add_heading('C Put the story in order', 1)
    doc.add_paragraph('Names to use: Schwann, Hooke, Virchow, Leeuwenhoek, Schleiden. Type the names from earliest to latest. Try without looking at page 1.')
    response(doc, 'timeline', 'Earliest → latest: ', 2)
    doc.add_heading('Complete the historical paragraph', 1)
    doc.add_paragraph('Word bank: cork  |  pond water  |  plants  |  animals  |  existing cells')
    sentences = [('Robert Hooke examined ', ' and named the tiny compartments cells.'), ('Anton van Leeuwenhoek observed tiny living organisms in ', '.'), ('Matthias Schleiden concluded that ', ' are made of cells.'), ('Theodor Schwann concluded that ', ' are made of cells.'), ('Rudolf Virchow helped establish that new cells come from ', '.')]
    for i, (before, after) in enumerate(sentences):
        p = doc.add_paragraph(before)
        field(p, f'paragraph_{i+1}')
        p.add_run(after)

    page(doc, 'D Choose an answer')
    doc.add_paragraph('Type A, B or C. You can revise your answer after reviewing the story.')
    for i, (prompt, choices, _) in enumerate(MC):
        p = doc.add_paragraph(f'{i+1}. {prompt}')
        p.paragraph_format.keep_with_next = True
        p.runs[0].bold = True
        doc.add_paragraph('   '.join(f'{letter}  {label}' for letter, label in zip('ABC', choices)))
        response(doc, f'multiple_choice_{i+1}', 'My answer: ')
    doc.add_heading('E Connect a discovery to a principle', 1)
    doc.add_paragraph('Write 1, 2 or 3 for the principle most directly supported. 1 = made of cells; 2 = basic unit of structure and function; 3 = cells from existing cells.')
    for i, statement in enumerate([
        'Plants and animals are made of cells.',
        'Cells build organisms and carry out basic life functions.',
        'A dividing cell produces new cells.',
    ]):
        response(doc, f'principle_match_{i+1}', statement + '  Principle: ')

    page(doc, 'F Explain the story in your own words')
    doc.add_paragraph('Use complete ideas. The response fields expand as you type. Compare your ideas with the separate answer key when you are ready.')
    for i, (prompt, _) in enumerate(TYPED):
        p = doc.add_paragraph(f'{i+1}. {prompt}')
        p.paragraph_format.keep_with_next = True
        p.runs[0].bold = True
        p = response(doc, f'written_response_{i+1}', lines=3)
        p.paragraph_format.space_after = Pt(11)
    path = OUT / 'cell-theory-worksheet.docx'
    doc.save(path)
    return path


def build_key():
    doc = document('Cell Theory and Scientists Answer Key', 'Use this separate key after attempting the worksheet. Accept equivalent wording that communicates the scientific idea. Short written responses are examples, not the only correct sentences.')
    doc.add_heading('A The three parts', 1)
    for statement in ['1. All living organisms are made of one or more cells.', '2. The cell is the basic unit of structure and function in living organisms.', '3. All cells come from pre-existing cells.']:
        doc.add_paragraph(statement)
    doc.add_heading('B Scientist matching', 1)
    doc.add_paragraph('1 C  Robert Hooke\n2 E  Anton van Leeuwenhoek\n3 D  Matthias Schleiden\n4 A  Theodor Schwann\n5 B  Rudolf Virchow')
    doc.add_heading('C Chronological order', 1)
    doc.add_paragraph('1665 Hooke → 1673 Leeuwenhoek → 1838 Schleiden → 1839 Schwann → 1858 Virchow')
    doc.add_paragraph('The story moves from cork to microscopic life, then plants and animals, and finally the origin of new cells.')
    doc.add_heading('Historical paragraph', 1)
    doc.add_paragraph('cork; pond water; plants; animals; existing cells')
    doc.add_heading('D Multiple choice', 1)
    doc.add_paragraph('1 A  Hooke\n2 B  Schleiden\n3 C  Schleiden and Schwann\n4 A  Virchow\n5 B  Cork contains tiny compartments')
    doc.add_heading('E Discovery to principle', 1)
    doc.add_paragraph('1 → Part 1; 2 → Part 2; 3 → Part 3. Focus on what organisms are made of, the role of the cell, and the origin of new cells.')
    page(doc, 'F Examples of written responses')
    titles = ['The three principles', 'How Hooke opened new questions', 'The plant and animal connection', 'The contribution of Virchow', 'Why microscopes mattered']
    for i, (_, answer) in enumerate(TYPED):
        doc.add_heading(f'{i+1} {titles[i]}', 2)
        doc.add_paragraph(answer)
    doc.add_heading('What to check', 1)
    doc.add_paragraph('Hooke saw dead cork cell walls, not living green cells. Leeuwenhoek revealed microscopic life. Schleiden studied plants; Schwann studied animals. Their work supported the first two principles. Virchow helped establish cells from cells; he did not invent cells.')
    doc.add_paragraph('The 1673 marker introduces Leeuwenhoek’s reports; his microorganism observations followed over several years. The 1858 marker identifies Virchow’s book, not the first observation of cell division.')
    path = OUT / 'cell-theory-answer-key.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    for file in [build_worksheet(), build_key()]:
        print(file)
